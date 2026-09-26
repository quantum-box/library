"""Write the hand-made MEXT fixtures used by `library food import` tests.

The layout copies the official 日本食品標準成分表（八訂）増補2023年 Excel
(sheet 表全体; header rows 2-12, component identifiers in row 12) with a
handful of columns and rows. The values are made up to cover every
notation the importer must keep apart: `-`, `0`, `Tr`, `(0)`, `(Tr)`,
`(12.3)`, a footnote dagger, an unknown `*`, leading-zero food numbers,
refuse rates and remarks. The errata copy the official 正誤表 sheets.

    python3 make_fixtures.py   # needs openpyxl
"""

from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).parent
SP = "　"

WIDTH = 19


def row(cells=None):
    r = [None] * WIDTH
    for col, value in (cells or {}).items():
        r[col] = value
    return r


def header_rows(offset=0):
    def shift(cells):
        return {c + offset: v for c, v in cells.items()}

    return [
        row(shift({16: "更新日：2026年3月27日"})),
        row(shift({0: f"食{SP}品{SP}群", 1: f"食{SP}品{SP}番{SP}号", 2: f"索{SP}引{SP}番{SP}号",
                   3: f"可{SP}{SP}食{SP}{SP}部{SP}{SP} 100{SP}{SP}g{SP}{SP}当{SP}{SP}た{SP}{SP}り", 17: f"備{SP}{SP}考"})),
        row(shift({3: f"食{SP}品{SP}名", 4: f"廃{SP}棄{SP}率", 5: "エネルギー", 7: f"たんぱく質", 9: "炭水化物",
                   13: "無 機 質", 15: f"ビ{SP}{SP}タ{SP}{SP}ミ{SP}{SP}ン"})),
        row(shift({7: "アミノ酸組成による\nたんぱく質", 8: "たんぱく質", 9: "利用可能炭水化物", 12: "食物繊維総量", 13: "ヨウ素",
                   14: "ナ ト リ ウ ム", 15: "ビタミンK", 16: "ビ\nタ\nミ\nン\nＢ１２"})),
        row(shift({9: "利用可能炭水化物\n（単糖当量）", 11: f"差引き法による\n利用可能炭水化物{SP}"})),
        row(),
        row(shift({3: "単位", 4: "%", 5: "kJ", 6: "kcal", 7: "(…………… g ………………)", 13: "(…… μg……)", 14: "mg",
                   15: "(…… μg……)"})),
        row(shift({3: "成分識別子", 4: "REFUSE", 5: "ENERC", 6: "ENERC_KCAL", 7: "PROTCAA", 8: "PROT-",
                   9: "CHOAVLM", 11: "CHOAVLDF-", 12: "FIB-", 13: "ID", 14: "NA", 15: "VITK ", 16: "VITB12"})),
    ]


FOODS = [
    # group, code, index, name, refuse, ENERC, KCAL, PROTCAA, PROT-, CHOAVLM, marker, CHOAVLDF-, FIB-, ID, NA, VITK, VITB12, remarks
    ["01", "01001", "0001", f"アマランサス{SP}玄穀", 0, 1452, 343, "(11.3)", 12.7, 63.5, "*", 59.9, 7.4, 1, 1, "(0)", "(0)", ""],
    ["06", "06153", "0582", f"（たまねぎ類）{SP}たまねぎ{SP}りん茎{SP}生", 6, 138, 33, 0.7, "1.0", "(3.9)", "*", 7.0, 1.5, 1, 2, 0, "-",
     "別名： 玉ねぎ、オニオン\n廃棄部位： 皮（保護葉）、底盤部及び頭部"],
    ["06", "06154", "0583", f"（たまねぎ類）{SP}たまねぎ{SP}りん茎{SP}ゆで", 0, 200, 31, 0.5, 0.8, "20.3†", None, 6.9, 1.7, 1, "(Tr)", "Tr", "(Tr)",
     "†は規定法による測定値"],
    ["10", "10330", "1520", f"＜魚類＞{SP}（あじ類）{SP}まあじ{SP}皮つき{SP}フライ", "0", 1134, 270, "(16.6)", 20.1, "(7.9)", None, 12.3, None, "*", 1, 7, None,
     "ヨウ素： 第3章参照"],
    ["01", "1002", "0002", f"アマランサス{SP}ゆで", 0, 1, 1, 1, 1, 1, None, 1, 1, 1, 1, 1, 1, ""],
    ["18", "99999", "9999", "重複した食品", 0, 1, 1, 1, 1, 1, None, 1, 1, 1, 1, 1, 1, ""],
    ["18", "99999", "9999", "重複した食品", 0, 1, 1, 1, 1, 1, None, 1, 1, 1, 1, 1, 1, ""],
    ["18", "18001", "2400", f"和風料理{SP}和え物類{SP}青菜の白和え", 120, 1, 1, 1, 1, 1, None, 1, 1, 1, 1, 1, 1, ""],
]


def main_table():
    wb = Workbook()
    ws = wb.active
    ws.title = "表全体"
    for r in header_rows():
        ws.append(r)
    for f in FOODS:
        ws.append(f)
    ws.append([None] * WIDTH)
    for name in ["1穀類", "6野菜類", "10魚介類", "18調理済み流通食品類"]:
        wb.create_sheet(name)
    wb.save(HERE / "mext_table_fixture.xlsx")


ERRATA_HEADER = ["変更対象", "頁", "食品番号", "索引番号", "食品名等", "項目等", "誤", "正", "備考"]
TARGET = "（八訂）増補2023年本表"


def errata():
    wb = Workbook()
    ch1 = wb.active
    ch1.title = "本表第1章"
    ch1.append([f"日本食品標準成分表（八訂）増補2023年{SP}正誤"])
    ch1.append([])
    ch1.append([None] * 8 + ["令和8年3月27日"])
    ch1.append(ERRATA_HEADER)
    ch1.append([TARGET, "表13", "10390", "-", "まあじ フライ", "卵液", 7.5, 9.2])

    ch2 = wb.create_sheet("本表第2章")
    ch2.append([f"日本食品標準成分表（八訂）増補2023年{SP}正誤"])
    ch2.append([])
    ch2.append([None] * 8 + ["令和8年3月27日"])
    ch2.append(ERRATA_HEADER)
    for e in [
        # already in the table: the official Excel was republished with it
        ["01001", "0001", "アマランサス 玄穀", "たんぱく質", "12.0", 12.7],
        # still wrong in the table: gets applied
        ["06153", "0582", "たまねぎ 生", f"エネルギー{SP}kcal", 33, 31],
        # remark fragment
        ["06153", "0582", "たまねぎ 生", "備考", "オニオン", "オニオン、玉葱"],
        # neither the wrong nor the right value: a conflict
        ["10330", "1520", "まあじ フライ", "ビタミンK", 5, "-"],
        # food name fragment, already applied
        ["10330", "1520", "まあじ フライ", "食品名", "皮付き", "皮つき"],
        # energy marker only, already applied
        ["01001", "0001", "アマランサス 玄穀", "利用可能炭水化物（単糖当量）\nアスタリスク", 63.5, "63.5*"],
        # an item the importer cannot place
        ["01001", "0001", "アマランサス 玄穀", "ほげ成分", 1, 2],
        # whole row: see sheet 本表
        ["06154", "0583", "たまねぎ ゆで", "各成分", "シート「本表」参照", None],
    ]:
        ch2.append([TARGET, None] + e)

    full = wb.create_sheet("本表")
    for r in header_rows(offset=1):
        full.append(r)
    full.cell(row=1, column=1, value=TARGET)
    wrong = ["誤"] + FOODS[2][:]
    right = ["正"] + FOODS[2][:]
    wrong[1 + 5] = 210  # ENERC: table has 200 → already applied
    wrong[1 + 6] = 30   # ENERC_KCAL: table has 31 → already applied
    right[1 + 14] = "(Tr)"
    wrong[1 + 14] = 3   # NA: table has (Tr) → already applied
    right[1 + 8] = 0.9  # PROT-: table has 0.8 = wrong → applied
    full.append(wrong)
    full.append(right)

    other = wb.create_sheet("ア第2章")
    other.append(ERRATA_HEADER)
    other.append([TARGET, None, "01001", "0001", "アマランサス", "イソロイシン", 1, 2])
    wb.save(HERE / "mext_errata_fixture.xlsx")


if __name__ == "__main__":
    main_table()
    errata()
